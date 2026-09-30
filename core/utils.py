from datetime import date
from functools import wraps

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import redirect
from django.utils import timezone


def role_required(*roles):
    """Admin anaruhusiwa kila mahali. Roles nyingine lazima zitajwe."""
    def decorator(view):
        @wraps(view)
        @login_required
        def wrapper(request, *args, **kwargs):
            user = request.user
            if user.is_admin or user.role in roles:
                return view(request, *args, **kwargs)
            messages.error(request, "Huna ruhusa ya kufanya hilo.")
            return redirect("dashboard")
        return wrapper
    return decorator


def branch_required(view):
    """Afisa au meneja bila tawi hawezi kusajili chochote: kila rekodi lazima iwe na tawi."""
    @wraps(view)
    @login_required
    def wrapper(request, *args, **kwargs):
        if not request.user.is_admin and not request.user.branch_id:
            messages.error(request, "Akaunti yako haina tawi. Mwombe admin akuwekee tawi kwanza.")
            return redirect("dashboard")
        return view(request, *args, **kwargs)
    return wrapper


admin_required = role_required()
manager_required = role_required("MANAGER")


def scope_by_branch(qs, user, field="branch"):
    """Admin anaona kila kitu, wengine wanaona tawi lao tu."""
    if user.is_admin:
        return qs
    if not user.branch_id:
        return qs.none()
    return qs.filter(**{f"{field}_id": user.branch_id})


def branch_filter(request, qs, field="branch"):
    """Kama scope_by_branch, pamoja na admin kuchagua tawi kwa ?tawi=ID."""
    qs = scope_by_branch(qs, request.user, field)
    tawi = request.GET.get("tawi", "")
    if request.user.is_admin and tawi.isdigit():
        qs = qs.filter(**{f"{field}_id": int(tawi)})
    return qs


def date_range(request, default_start=None):
    today = timezone.localdate()

    def parse(name, default):
        try:
            return date.fromisoformat(request.GET.get(name, ""))
        except ValueError:
            return default

    start = parse("kuanzia", default_start or today.replace(day=1))
    end = parse("mpaka", today)
    if start > end:
        start, end = end, start
    return start, end


def total(qs, field="amount"):
    return qs.aggregate(t=Sum(field))["t"] or 0


def paginate(request, qs, per_page=25):
    return Paginator(qs, per_page).get_page(request.GET.get("page"))


def wants_excel(request):
    return request.GET.get("export") in ("xlsx", "excel", "csv")


class Sheet:
    """Karatasi moja ya Excel: kichwa, safu, mistari na (hiari) jumla ya chini."""
    def __init__(self, name, header, rows, totals=None, title=""):
        self.name, self.header, self.rows, self.totals, self.title = name[:31], header, rows, totals, title


def excel_response(filename, sheets, subtitle=""):
    """
    Faili la Excel (.xlsx) lenye muonekano wa kampuni: kichwa cheusi na dhahabu, namba zenye
    koma (zinazoweza kujumlishwa kwenye Excel), vichujio, na safu ya juu inayobaki ukishuka.
    """
    from decimal import Decimal
    from datetime import date as _date, datetime as _dt

    from django.conf import settings
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if isinstance(sheets, Sheet):
        sheets = [sheets]
    wb = Workbook()
    wb.remove(wb.active)
    ink, gold = PatternFill("solid", fgColor="111111"), "F2C24E"
    for sh in sheets:
        ws = wb.create_sheet(sh.name)
        ncols = max(1, len(sh.header))
        ws.append([settings.COMPANY_NAME])
        ws["A1"].font = Font(bold=True, size=14, color="8A5F00")
        ws.append([sh.title or sh.name])
        ws["A2"].font = Font(bold=True, size=12)
        ws.append([subtitle or f"Imetolewa {timezone.localtime():%d/%m/%Y %H:%M}"])
        ws["A3"].font = Font(italic=True, color="777777")
        ws.append([])
        ws.append(list(sh.header))
        head_row = ws.max_row
        for c in range(1, ncols + 1):
            cell = ws.cell(row=head_row, column=c)
            cell.fill, cell.font = ink, Font(bold=True, color=gold)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [len(str(h)) + 2 for h in sh.header]
        for row in list(sh.rows) + ([sh.totals] if sh.totals else []):
            out = []
            for v in row:
                if isinstance(v, Decimal):
                    v = float(v)
                if isinstance(v, _dt):
                    v = timezone.localtime(v).replace(tzinfo=None) if timezone.is_aware(v) else v
                out.append(v)
            ws.append(out)
            r = ws.max_row
            for i, v in enumerate(out, start=1):
                cell = ws.cell(row=r, column=i)
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    cell.number_format = "#,##0"
                elif isinstance(v, _dt):
                    cell.number_format = "dd/mm/yyyy hh:mm"
                elif isinstance(v, _date):
                    cell.number_format = "dd/mm/yyyy"
                if i <= len(widths):
                    widths[i - 1] = max(widths[i - 1], min(48, len(str(v if v is not None else "")) + 2))
        if sh.totals:
            for c in range(1, ncols + 1):
                ws.cell(row=ws.max_row, column=c).font = Font(bold=True)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, w)
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
        last = ws.max_row - (1 if sh.totals else 0)
        if last > head_row:
            ws.auto_filter.ref = f"A{head_row}:{get_column_letter(ncols)}{last}"
    if not filename.endswith(".xlsx"):
        filename = filename.rsplit(".", 1)[0] + ".xlsx"
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
