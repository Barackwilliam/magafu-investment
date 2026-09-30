from datetime import timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from core.models import Branch
from customers.models import Customer
from finance.models import CashEntry

from . import services
from .models import Loan, LoanProduct


class Fixtures(TestCase):
    def setUp(self):
        self.b1 = Branch.objects.create(name="Kariakoo")
        self.b2 = Branch.objects.create(name="Mbezi")
        self.admin = User.objects.create_user("admin", password="pass1234", role="ADMIN", first_name="Admin")
        self.manager = User.objects.create_user("meneja", password="pass1234", role="MANAGER", branch=self.b1)
        self.officer = User.objects.create_user("afisa", password="pass1234", role="OFFICER", branch=self.b1)
        self.product = LoanProduct.objects.create(name="Siku 30", interest_rate=20, duration_days=30,
                                                  form_fee=10000, penalty_per_day=1000)
        self.c1 = Customer.objects.create(branch=self.b1, first_name="Hawa", last_name="Musa", gender="F", phone="0755000001")
        self.c2 = Customer.objects.create(branch=self.b2, first_name="Juma", last_name="Ali", gender="M", phone="0755000002")

    def _active_loan(self, customer=None, amount=200000):
        loan = services.create_loan(customer=customer or self.c1, product=self.product,
                                    principal=Decimal(amount), user=self.officer, form_fee_paid=True)
        services.approve_loan(loan.pk, self.manager)
        return services.disburse_loan(loan.pk, self.manager)


class LoanFlowTests(Fixtures):

    def test_interest_and_full_flow(self):
        loan = self._active_loan()
        self.assertEqual(loan.total_payable, Decimal(240000))
        self.assertEqual(loan.due_date, timezone.localdate() + timedelta(days=30))
        self.assertEqual(CashEntry.objects.filter(entry_type="FORM_FEE").count(), 1)
        services.record_payment(loan.pk, amount=Decimal(100000), user=self.officer)
        loan.refresh_from_db()
        self.assertEqual(loan.balance, Decimal(140000))
        with self.assertRaises(services.LoanError):
            services.record_payment(loan.pk, amount=Decimal(150000), user=self.officer)
        services.record_payment(loan.pk, amount=Decimal(140000), user=self.officer)
        loan.refresh_from_db()
        self.assertEqual(loan.status, Loan.Status.COMPLETED)

    def test_one_open_loan_per_customer(self):
        self._active_loan()
        with self.assertRaises(services.LoanError):
            services.create_loan(customer=self.c1, product=self.product, principal=Decimal(1000), user=self.officer)

    def test_product_change_does_not_affect_old_loans(self):
        loan = self._active_loan()
        self.product.interest_rate = 30
        self.product.save()
        loan.refresh_from_db()
        self.assertEqual(loan.total_payable, Decimal(240000))

    def test_overdue_and_auto_penalty_once_per_day(self):
        loan = self._active_loan()
        Loan.objects.filter(pk=loan.pk).update(due_date=timezone.localdate() - timedelta(days=3))
        self.assertEqual(services.apply_daily_penalties(), 1)
        self.assertEqual(services.apply_daily_penalties(), 0)
        loan = Loan.objects.with_totals().get(pk=loan.pk)
        self.assertTrue(loan.is_overdue)
        self.assertEqual(loan.balance_total, Decimal(241000))

    def test_branch_isolation(self):
        other = self._active_loan(customer=self.c2)
        self.client.login(username="afisa", password="pass1234")
        self.assertEqual(self.client.get(reverse("loan_detail", args=[other.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("customer_detail", args=[self.c2.pk])).status_code, 404)

    def test_officer_cannot_approve(self):
        loan = services.create_loan(customer=self.c1, product=self.product, principal=Decimal(50000), user=self.officer)
        self.client.login(username="afisa", password="pass1234")
        self.client.post(reverse("loan_approve", args=[loan.pk]))
        loan.refresh_from_db()
        self.assertEqual(loan.status, Loan.Status.PENDING)
        self.assertEqual(self.client.get(reverse("product_list")).status_code, 302)

    def test_pay_via_view(self):
        loan = self._active_loan()
        self.client.login(username="afisa", password="pass1234")
        r = self.client.post(reverse("loan_pay", args=[loan.pk]), {"amount": "40000", "method": "MPESA", "reference": "QX1"})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(loan.repayments.count(), 1)

    def test_all_pages_render(self):
        loan = self._active_loan()
        services.record_payment(loan.pk, amount=Decimal(20000), user=self.officer)
        CashEntry.objects.create(branch=self.b1, entry_type="EXPENSE", amount=5000, recorded_by=self.officer)
        pending = services.create_loan(customer=self.c2, product=self.product, principal=Decimal(50000), user=self.admin)
        urls = [
            reverse("dashboard"), reverse("customer_list"), reverse("customer_create"),
            reverse("customer_detail", args=[self.c1.pk]), reverse("customer_edit", args=[self.c1.pk]),
            reverse("loan_list"), reverse("loan_list") + "?hali=sugu", reverse("loan_list") + "?hali=inasubiri",
            reverse("loan_apply"), reverse("loan_detail", args=[loan.pk]), reverse("loan_detail", args=[pending.pk]),
            reverse("loan_statement", args=[loan.pk]), reverse("product_list"), reverse("product_create"),
            reverse("cash_list"), reverse("cash_create"), reverse("report_daily"), reverse("report_collections"),
            reverse("report_collections") + "?export=csv", reverse("report_disbursements"),
            reverse("report_overdue"), reverse("report_overdue") + "?export=csv", reverse("report_branches"),
            reverse("sms_log"), reverse("branch_list"), reverse("branch_create"), reverse("user_list"),
            reverse("user_create"), reverse("password_change"),
        ]
        for username in ("admin", "meneja", "afisa"):
            self.client.login(username=username, password="pass1234")
            for url in urls:
                r = self.client.get(url)
                allowed = (200, 302)
                if url == reverse("loan_detail", args=[pending.pk]) and username != "admin":
                    allowed = (404,)  # mkopo wa tawi jingine
                self.assertIn(r.status_code, allowed, f"{username} {url} -> {r.status_code}")
            self.client.logout()
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 302)
        self.assertEqual(self.client.get(reverse("login")).status_code, 200)

    def test_cards_and_expected_collection(self):
        from core.stats import dashboard_cards
        loan = self._active_loan()  # 200,000 + 20% = 240,000, siku 30 kila siku -> 8,000 kwa siku
        Loan.objects.filter(pk=loan.pk).update(disbursed_at=timezone.now() - timedelta(days=5))
        services.record_payment(loan.pk, amount=Decimal(3000), user=self.officer)
        c = dashboard_cards(self.admin)
        self.assertEqual(c["makusanyo"]["makadirio"], Decimal(8000))
        self.assertEqual(c["makusanyo"]["leo"], Decimal(3000))
        self.assertEqual(c["makusanyo"]["hazijakusanywa"], Decimal(5000))
        self.assertEqual(c["makusanyo"]["asilimia"], 38)
        self.assertEqual(c["mikopo"]["wadaiwa"], Decimal(237000))
        self.assertEqual(c["mapato"]["fomu"], Decimal(10000))
        self.assertEqual(c["wateja"]["wanadaiwa"], 1)
        self.assertEqual(c["maombi"]["jumla"], 1)
        # Afisa wa tawi jingine haoni takwimu za tawi hili
        other = User.objects.create_user("mbezi", password="x", role="OFFICER", branch=self.b2)
        self.assertEqual(dashboard_cards(other)["mikopo"]["jumla"], Decimal(0))

    def test_branch_list_actions(self):
        self.client.login(username="admin", password="pass1234")
        r = self.client.post(reverse("branch_create"), {"name": "Mwenge", "phone": "0711000000", "is_active": "on"})
        self.assertEqual(r.status_code, 302)
        mwenge = Branch.objects.get(name="Mwenge")
        self.assertEqual(mwenge.created_by, self.admin)
        page = self.client.get(reverse("branch_list")).content.decode()
        for col in ("S/N", "Msajili", "Tarehe", "Actions"):
            self.assertIn(col, page)
        self.client.post(reverse("branch_toggle", args=[mwenge.pk]))
        mwenge.refresh_from_db()
        self.assertFalse(mwenge.is_active)
        # Tawi lililozimwa halionekani kwenye fomu ya mteja mpya
        self.assertNotIn(">Mwenge</option>", self.client.get(reverse("customer_create")).content.decode())
        # Tawi lenye wateja haliwezi kufutwa, tupu linafutika
        self.client.post(reverse("branch_delete", args=[self.b1.pk]))
        self.assertTrue(Branch.objects.filter(pk=self.b1.pk).exists())
        self.client.post(reverse("branch_delete", args=[mwenge.pk]))
        self.assertFalse(Branch.objects.filter(pk=mwenge.pk).exists())


class BugFixTests(Fixtures):
    """Kila test hapa inalinda bug iliyokuwepo kweli."""

    def test_cash_form_shows_only_own_branch_customers(self):
        """BUG: afisa aliona wateja wa matawi yote, admin hakuona yeyote."""
        from finance.forms import CashEntryForm
        officer_ids = set(CashEntryForm(user=self.officer).fields["customer"].queryset.values_list("pk", flat=True))
        admin_ids = set(CashEntryForm(user=self.admin).fields["customer"].queryset.values_list("pk", flat=True))
        self.assertEqual(officer_ids, {self.c1.pk})
        self.assertEqual(admin_ids, {self.c1.pk, self.c2.pk})

    def test_cash_customer_must_match_branch(self):
        from finance.forms import CashEntryForm
        form = CashEntryForm({"entry_type": "FORM_FEE", "date": timezone.localdate(), "amount": "5000",
                              "branch": self.b1.pk, "customer": self.c2.pk}, user=self.admin)
        self.assertFalse(form.is_valid())
        self.assertIn("customer", form.errors)

    def test_staff_without_branch_is_stopped_not_crashed(self):
        """BUG: afisa bila tawi alipata server error (branch=None) akisajili mteja."""
        User.objects.create_user("bila", password="pass1234", role="OFFICER")
        self.client.login(username="bila", password="pass1234")
        for name in ("customer_create", "cash_create", "loan_apply"):
            r = self.client.post(reverse(name), {"first_name": "X", "amount": "100"})
            self.assertRedirects(r, reverse("dashboard"), fetch_redirect_response=False)
        self.assertFalse(Customer.objects.filter(first_name="X").exists())

    def test_admin_cannot_lock_themselves_out(self):
        self.client.login(username="admin", password="pass1234")
        self.client.post(reverse("user_edit", args=[self.admin.pk]), {
            "first_name": "Admin", "username": "admin", "role": "OFFICER", "branch": self.b1.pk,
        })
        self.admin.refresh_from_db()
        self.assertEqual(self.admin.role, "ADMIN")
        self.assertTrue(self.admin.is_active)

    def test_product_limits_are_validated(self):
        from .forms import LoanProductForm
        form = LoanProductForm({"name": "Mbaya", "interest_rate": "150", "duration_days": "0",
                                "repayment_frequency": "DAILY", "form_fee": "0", "penalty_per_day": "0",
                                "min_amount": "500000", "max_amount": "1000"})
        self.assertFalse(form.is_valid())
        for field in ("interest_rate", "duration_days", "max_amount"):
            self.assertIn(field, form.errors)

    def test_customer_with_open_loan_cannot_change_branch(self):
        from customers.forms import CustomerForm
        self._active_loan()
        data = {"first_name": "Hawa", "last_name": "Musa", "gender": "F", "phone": "0755000001",
                "branch": self.b2.pk, "is_active": "on"}
        form = CustomerForm(data, instance=self.c1, user=self.admin)
        self.assertFalse(form.is_valid())
        self.assertIn("branch", form.errors)

    def test_phone_numbers_are_normalised(self):
        from customers.forms import CustomerForm
        form = CustomerForm({"first_name": "A", "last_name": "B", "gender": "M", "phone": "+255 712 345 678",
                             "guarantor_phone": "712 000 111", "branch": self.b1.pk}, user=self.admin)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["phone"], "0712345678")
        self.assertEqual(form.cleaned_data["guarantor_phone"], "0712000111")

    def test_penalties_apply_without_cron(self):
        """BUG: bila Cron Job (inayolipiwa Render) faini hazikuwekwa kamwe."""
        from django.core.cache import cache
        cache.clear()
        loan = self._active_loan()
        Loan.objects.filter(pk=loan.pk).update(due_date=timezone.localdate() - timedelta(days=2))
        self.client.login(username="afisa", password="pass1234")
        self.client.get(reverse("dashboard"))
        self.client.get(reverse("dashboard"))
        self.assertEqual(loan.penalty_entries.count(), 1)


class DeployTests(TestCase):
    def test_production_static_files_build(self):
        """BUG: faili la JS lilitaja source map isiyokuwepo, na collectstatic ya Render ingeshindwa."""
        import tempfile
        from django.core.management import call_command
        from django.test import override_settings
        with tempfile.TemporaryDirectory() as tmp, override_settings(
            STATIC_ROOT=tmp,
            STORAGES={"default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
                      "staticfiles": {"BACKEND": "whitenoise.storage.CompressedManifestStaticFilesStorage"}},
        ):
            call_command("collectstatic", interactive=False, verbosity=0)

    def test_search_finds_customer_and_loan_number(self):
        from core.models import Branch
        b = Branch.objects.create(name="Tawi")
        User.objects.create_user("a", password="pass1234", role="ADMIN")
        c = Customer.objects.create(branch=b, first_name="Rehema", last_name="Lema", gender="F", phone="0712000999")
        self.client.login(username="a", password="pass1234")
        r = self.client.get(reverse("search"), {"q": "+255 712 000 999"})
        self.assertRedirects(r, reverse("customer_detail", args=[c.pk]))


class CustomerPagePaymentTests(Fixtures):
    def test_pay_from_customer_page_returns_there(self):
        """Mtumiaji anarekodi malipo kutoka ukurasa wa mteja, bila kwenda ukurasa wa mkopo."""
        loan = self._active_loan()
        self.client.login(username="afisa", password="pass1234")
        page = self.client.get(reverse("customer_detail", args=[self.c1.pk])).content.decode()
        self.assertIn("Rekodi malipo aliyoleta", page)
        back = reverse("customer_detail", args=[self.c1.pk]) + "#lipa"
        r = self.client.post(reverse("loan_pay", args=[loan.pk]), {"amount": "5000", "method": "CASH", "next": back})
        self.assertRedirects(r, back, fetch_redirect_response=False)
        loan = Loan.objects.with_totals().get(pk=loan.pk)
        self.assertEqual(loan.paid_total, Decimal(5000))
        self.assertIn("5,000", self.client.get(back).content.decode())

    def test_next_cannot_redirect_off_site(self):
        loan = self._active_loan()
        self.client.login(username="afisa", password="pass1234")
        r = self.client.post(reverse("loan_pay", args=[loan.pk]), {"amount": "5000", "method": "CASH", "next": "https://evil.example/"})
        self.assertRedirects(r, reverse("loan_detail", args=[loan.pk]), fetch_redirect_response=False)


class OneStepLoanTests(Fixtures):
    def test_toa_mkopo_is_active_immediately(self):
        """Mtumiaji mmoja: 'Toa mkopo' inaanza deni papo hapo, bila kuthibitisha na kutoa pesa kando."""
        self.client.login(username="admin", password="pass1234")
        r = self.client.post(reverse("loan_apply"), {"customer": self.c1.pk, "product": self.product.pk,
                                                     "principal": "100000", "form_fee_paid": "on"})
        loan = Loan.objects.get(customer=self.c1)
        self.assertEqual(loan.status, Loan.Status.ACTIVE)
        self.assertIsNotNone(loan.due_date)
        self.assertRedirects(r, reverse("customer_detail", args=[self.c1.pk]) + "#lipa", fetch_redirect_response=False)
        self.assertIn("Rekodi malipo aliyoleta", self.client.get(reverse("customer_detail", args=[self.c1.pk])).content.decode())

    def test_old_pending_loan_can_be_given_in_one_click(self):
        loan = services.create_loan(customer=self.c1, product=self.product, principal=Decimal(50000), user=self.officer)
        self.client.login(username="admin", password="pass1234")
        self.client.post(reverse("loan_disburse", args=[loan.pk]))
        loan.refresh_from_db()
        self.assertEqual(loan.status, Loan.Status.ACTIVE)


class ExcelExportTests(Fixtures):
    def test_every_report_downloads_a_real_excel_file(self):
        from io import BytesIO
        from openpyxl import load_workbook
        loan = self._active_loan()
        services.record_payment(loan.pk, amount=Decimal(20000), user=self.officer)
        CashEntry.objects.create(branch=self.b1, entry_type="EXPENSE", amount=5000, recorded_by=self.officer)
        urls = ["customer_list", "loan_list", "cash_list", "report_daily", "report_collections",
                "report_disbursements", "report_overdue", "report_branches"]
        self.client.login(username="admin", password="pass1234")
        for name in urls:
            r = self.client.get(reverse(name), {"export": "xlsx"})
            self.assertEqual(r["Content-Type"],
                             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", name)
            self.assertIn(".xlsx", r["Content-Disposition"], name)
            wb = load_workbook(BytesIO(r.content))
            self.assertTrue(wb.sheetnames, name)
        # namba ni namba halisi (zinajumlishika), si maandishi
        wb = load_workbook(BytesIO(self.client.get(reverse("report_collections"), {"export": "xlsx"}).content))
        ws = wb.active
        amounts = [row[4] for row in ws.iter_rows(min_row=6, values_only=True) if row[1]]
        self.assertIn(20000, amounts)

    def test_officer_export_only_contains_own_branch(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.client.login(username="afisa", password="pass1234")
        wb = load_workbook(BytesIO(self.client.get(reverse("customer_list"), {"export": "xlsx"}).content))
        names = [row[0] for row in wb.active.iter_rows(min_row=6, values_only=True)]
        self.assertIn("Hawa Musa", names)
        self.assertNotIn("Juma Ali", names)
